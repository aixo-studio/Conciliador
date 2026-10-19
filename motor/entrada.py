"""Paso 1 del proceso: leer la entrada.

Busca en la salida de la etapa 1 ("conciliacion v3") los dos archivos de un
banco, controla que cumplan el contrato (ver CLAUDE.md) y los lee. Si algo no
cumple, frena y dice qué archivo, qué fila y por qué. Nunca modifica la entrada.
"""
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
from typing import Optional

import openpyxl

# La carpeta de datos está al lado de la carpeta del repo.
CARPETA_DATOS = Path(__file__).resolve().parent.parent.parent / "conciliacion resultados"
ENTRADA = CARPETA_DATOS / "v3" / "output"   # lo que deja la etapa 1
SALIDA = CARPETA_DATOS / "v4" / "output"    # lo que deja este repo

# Lo que pide el contrato con la etapa 1.
HOJA = "Movimientos"
ENCABEZADO = ("Fecha", "Concepto", "Descripción", "Monto")

# Los estados posibles de una fila.
SIN_CRUZAR = "Sin cruzar"
CRUZADO = "Cruzado"


@dataclass
class Movimiento:
    """Una fila de la entrada. Al leerla queda "Sin cruzar"; el estado, el número
    de cruce y el criterio los completa el cruce.
    """
    fila: int                            # número de fila en el Excel de entrada
    fecha: date
    concepto: str
    descripcion: str
    monto_centavos: int                  # el monto en centavos enteros
    estado: str = SIN_CRUZAR
    numero_cruce: Optional[int] = None   # une las filas del banco y del sistema que cruzaron
    criterio: str = ""                   # cómo se cruzó (hoy, solo "Exacto")


def es_excel(ruta: Path) -> bool:
    """Recibe una ruta. Devuelve True si es un .xlsx de datos. Los archivos
    ocultos y los temporales que deja Excel abierto ("~$...") no cuentan.
    """
    return (ruta.is_file() and ruta.suffix.lower() == ".xlsx"
            and not ruta.name.startswith(("~$", ".")))


def encontrar_archivos(carpeta: Path) -> tuple[Path, Path]:
    """Recibe la carpeta de un banco en la salida de v3.
    Devuelve el archivo del banco y el del sistema, que se reconocen por el
    principio del nombre ("banco_..." y "sistema_..."). Frena si no hay
    exactamente uno de cada uno, o si hay un Excel que no es ninguno de los dos.
    """
    if not carpeta.is_dir():
        raise ValueError(f"no existe la carpeta {carpeta}")
    archivos = sorted(ruta for ruta in carpeta.iterdir() if es_excel(ruta))
    del_banco = [ruta for ruta in archivos if ruta.name.startswith("banco_")]
    del_sistema = [ruta for ruta in archivos if ruta.name.startswith("sistema_")]
    otros = [ruta.name for ruta in archivos if ruta not in del_banco + del_sistema]
    if otros:
        raise ValueError("hay archivos que no son ni del banco ni del sistema: " + ", ".join(otros))
    if (len(del_banco), len(del_sistema)) != (1, 1):
        raise ValueError(f"tiene que haber 1 archivo del banco y 1 del sistema; "
                         f"hay {len(del_banco)} y {len(del_sistema)}.")
    return del_banco[0], del_sistema[0]


def esta_vacia(valores) -> bool:
    """Recibe los valores de una fila. Devuelve True si no tiene ningún dato."""
    return all(valor in (None, "") for valor in valores)


def a_centavos(valor, fila: int) -> int:
    """Recibe el Monto de una fila. Lo devuelve en centavos enteros, redondeando:
    un decimal leído de Excel puede venir como 1500000.2499999 en vez de .25.
    Frena si no es un número.
    """
    if isinstance(valor, bool) or not isinstance(valor, (int, float)):
        raise ValueError(f"fila {fila}: el Monto no es un número ({valor!r}).")
    return round(valor * 100)


def leer_movimientos(ruta: Path) -> list[Movimiento]:
    """Recibe un archivo de la entrada. Controla la hoja y el encabezado y lee
    fila por fila. Devuelve un Movimiento por fila, en el orden del archivo.
    Frena si algo no cumple el contrato.
    """
    libro = openpyxl.load_workbook(ruta, read_only=True, data_only=True)
    try:
        if libro.sheetnames != [HOJA]:
            raise ValueError(f'tiene que tener una sola hoja, "{HOJA}"; '
                             f'tiene: {", ".join(libro.sheetnames)}.')
        filas = [tuple(valores) for valores in libro[HOJA].iter_rows(values_only=True)]
    finally:
        libro.close()
    # Fila 1: el encabezado, exacto y sin columnas de más.
    if not filas or filas[0][:4] != ENCABEZADO or not esta_vacia(filas[0][4:]):
        raise ValueError("el encabezado tiene que ser Fecha · Concepto · Descripción · Monto.")
    movimientos = []
    for numero, valores in enumerate(filas[1:], start=2):
        # Una fila vacía o con datos de más no se saltea: se frena, para no perder nada.
        if esta_vacia(valores):
            raise ValueError(f"fila {numero}: está vacía.")
        if not esta_vacia(valores[4:]):
            raise ValueError(f"fila {numero}: tiene datos fuera de las 4 columnas.")
        fecha, concepto, descripcion, monto = (valores + (None,) * 4)[:4]
        if isinstance(fecha, datetime):
            fecha = fecha.date()
        if not isinstance(fecha, date):
            raise ValueError(f"fila {numero}: la Fecha no es una fecha ({fecha!r}).")
        movimientos.append(Movimiento(
            fila=numero, fecha=fecha,
            concepto="" if concepto is None else str(concepto),
            descripcion="" if descripcion is None else str(descripcion),
            monto_centavos=a_centavos(monto, numero)))
    if not movimientos:
        raise ValueError("no tiene movimientos.")
    return movimientos


def resumir(movimientos: list[Movimiento]) -> dict:
    """Recibe los movimientos de un archivo. Devuelve cuántas filas son, cuánto
    suman (en centavos) y la primera y la última fecha.
    """
    fechas = [movimiento.fecha for movimiento in movimientos]
    return {"filas": len(movimientos),
            "suma": sum(movimiento.monto_centavos for movimiento in movimientos),
            "desde": min(fechas), "hasta": max(fechas)}
