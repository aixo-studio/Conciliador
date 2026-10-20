"""Paso 3 del proceso: controlar que no se perdió nada y escribir el Excel.

El Excel tiene tres hojas: "Resumen" (los controles), "Banco" y "Sistema"
(todas las filas de la entrada, en el mismo orden, con su estado).
"""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font

from motor.entrada import CRUZADO, SIN_CRUZAR, Movimiento

COLUMNAS = ("Fecha", "Código", "Descripción", "Monto", "Estado")


def contar(movimientos: list[Movimiento]) -> dict:
    """Recibe los movimientos de un lado. Devuelve, para el total, lo cruzado y
    lo sin cruzar: cuántas filas son y cuánto suman (en centavos).
    """
    grupos = {"Entrada": movimientos,
              CRUZADO: [movimiento for movimiento in movimientos if movimiento.estado == CRUZADO],
              SIN_CRUZAR: [movimiento for movimiento in movimientos if movimiento.estado == SIN_CRUZAR]}
    return {nombre: (len(grupo), sum(movimiento.monto_centavos for movimiento in grupo))
            for nombre, grupo in grupos.items()}


def controlar(banco: list[Movimiento], sistema: list[Movimiento]) -> dict:
    """Recibe los dos lados ya cruzados y controla que no se perdió nada:
    - en cada lado, cruzado + sin cruzar da la entrada, en filas y en suma;
    - lo cruzado suma lo mismo en el banco y en el sistema;
    - las parejas se apuntan entre sí: si la fila 41 del banco dice que cruzó
      con la 512 del sistema, la 512 del sistema tiene que decir la 41.
    Devuelve los conteos de cada lado. Si algo no cierra, frena.
    """
    conteos = {"Banco": contar(banco), "Sistema": contar(sistema)}
    for lado, conteo in conteos.items():
        filas = conteo[CRUZADO][0] + conteo[SIN_CRUZAR][0]
        suma = conteo[CRUZADO][1] + conteo[SIN_CRUZAR][1]
        if (filas, suma) != conteo["Entrada"]:
            raise ValueError(f"{lado}: cruzado + sin cruzar no da la entrada "
                             f"({filas} filas contra {conteo['Entrada'][0]}).")
    if conteos["Banco"][CRUZADO][1] != conteos["Sistema"][CRUZADO][1]:
        raise ValueError("lo cruzado no suma lo mismo en el banco y en el sistema.")
    for lado, propios, otros in (("Banco", banco, sistema), ("Sistema", sistema, banco)):
        controlar_parejas(lado, propios, {movimiento.fila: movimiento for movimiento in otros})
    return conteos


def controlar_parejas(lado: str, propios: list[Movimiento], otros_por_fila: dict) -> None:
    """Recibe un lado, sus movimientos y los del otro lado (por número de fila).
    Controla que cada fila cruzada apunte a filas cruzadas del otro lado que la
    apunten de vuelta, y que las sin cruzar no apunten a ninguna. Si no, frena.
    """
    for movimiento in propios:
        if (movimiento.estado == CRUZADO) != bool(movimiento.filas_cruzadas):
            raise ValueError(f"{lado} fila {movimiento.fila}: el estado no coincide con sus filas cruzadas.")
        for fila in movimiento.filas_cruzadas:
            pareja = otros_por_fila.get(fila)
            if pareja is None or pareja.estado != CRUZADO or movimiento.fila not in pareja.filas_cruzadas:
                raise ValueError(f"{lado} fila {movimiento.fila}: dice que cruzó con la fila {fila} "
                                 f"del otro lado, pero esa fila no la apunta de vuelta.")


def mostrar_filas(filas: list[int]):
    """Recibe las filas con las que cruzó un movimiento. Devuelve lo que va en la
    celda: vacía si no cruzó, el número si es una, "41, 42, 43" si son varias.
    """
    if not filas:
        return None
    return filas[0] if len(filas) == 1 else ", ".join(str(fila) for fila in filas)


def escribir_hoja(libro: Workbook, titulo: str, encabezado, filas, formatos: dict, anchos) -> None:
    """Recibe el libro, el nombre de la hoja, el encabezado, las filas, el formato
    de cada columna (letra → formato) y los anchos. Agrega la hoja al libro.
    """
    hoja = libro.create_sheet(titulo)
    hoja.append(list(encabezado))
    for celda in hoja[1]:
        celda.font = Font(bold=True)
    for fila in filas:
        hoja.append(list(fila))
    for letra, formato in formatos.items():
        for (celda,) in hoja.iter_rows(min_row=2, min_col=ord(letra) - 64, max_col=ord(letra) - 64):
            celda.number_format = formato
    for numero, ancho in enumerate(anchos):
        hoja.column_dimensions[chr(65 + numero)].width = ancho
    hoja.freeze_panes = "A2"  # el encabezado queda fijo al bajar


def escribir_excel(banco: list[Movimiento], sistema: list[Movimiento], conteos: dict, ruta: Path) -> None:
    """Recibe los dos lados cruzados, sus conteos y la ruta donde guardar.
    Escribe el Excel con las hojas "Resumen", "Banco" y "Sistema".
    """
    libro = Workbook()
    libro.remove(libro.active)  # se arranca sin la hoja vacía que trae por defecto
    resumen = [(lado, nombre, filas, suma / 100)
               for lado, conteo in conteos.items() for nombre, (filas, suma) in conteo.items()]
    escribir_hoja(libro, "Resumen", ("Lado", "Qué", "Filas", "Suma"), resumen,
                  {"D": "#,##0.00"}, (10, 12, 10, 20))
    # En cada hoja, la columna de la pareja dice en qué fila del otro documento está.
    for titulo, otro, movimientos in (("Banco", "Sistema", banco), ("Sistema", "Banco", sistema)):
        filas = [(m.fecha, m.codigo, m.descripcion, m.monto_centavos / 100,
                  m.estado, mostrar_filas(m.filas_cruzadas), m.criterio) for m in movimientos]
        escribir_hoja(libro, titulo, COLUMNAS + (f"Fila en {otro}", "Criterio"), filas,
                      {"A": "DD/MM/YYYY", "D": "#,##0.00"}, (12, 34, 50, 18, 12, 15, 10))
    libro.save(ruta)
