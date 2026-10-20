"""Prueba de los controles y del Excel de salida (paso 3), y de la corrida
completa de un banco, con archivos inventados.
"""
import openpyxl
import pytest

from cruzar import cruzar_banco
from motor.cruce_exacto import cruzar_exacto
from motor.salida import controlar, mostrar_filas
from pruebas.datos import carpeta_banco, dia
from pruebas.test_cruce_exacto import movimiento


def test_controla_entrada_igual_a_cruzado_mas_sin_cruzar():
    banco = [movimiento(2, 1, 100), movimiento(3, 2, -50)]
    sistema = [movimiento(2, 1, 100)]
    cruzar_exacto(banco, sistema)
    conteos = controlar(banco, sistema)
    assert conteos["Banco"] == {"Entrada": (2, 50), "Cruzado": (1, 100), "Sin cruzar": (1, -50)}
    assert conteos["Sistema"] == {"Entrada": (1, 100), "Cruzado": (1, 100), "Sin cruzar": (0, 0)}


def test_frena_si_una_fila_queda_con_un_estado_desconocido():
    banco, sistema = [movimiento(2, 1, 100)], [movimiento(2, 1, 100)]
    banco[0].estado = "Otro"
    with pytest.raises(ValueError, match="no da la entrada"):
        controlar(banco, sistema)


def test_frena_si_las_parejas_no_se_apuntan_entre_si():
    banco, sistema = [movimiento(2, 1, 100)], [movimiento(2, 1, 100)]
    cruzar_exacto(banco, sistema)
    sistema[0].filas_cruzadas = [7]
    with pytest.raises(ValueError, match="no la apunta de vuelta"):
        controlar(banco, sistema)


def test_frena_si_una_fila_sin_cruzar_apunta_a_otra():
    banco, sistema = [movimiento(2, 1, 100)], [movimiento(2, 1, 200)]
    banco[0].filas_cruzadas = [2]
    with pytest.raises(ValueError, match="el estado no coincide"):
        controlar(banco, sistema)


def test_varias_filas_se_muestran_separadas_por_coma():
    assert (mostrar_filas([]), mostrar_filas([512]), mostrar_filas([41, 42, 43])) == (None, 512, "41, 42, 43")


def test_cruza_un_banco_y_escribe_el_excel(tmp_path):
    entrada = carpeta_banco(tmp_path / "entrada", [
        [dia(1), "COMISION", "", -10.5], [dia(1), "TRANSFERENCIA", "", -1500]],
        [[dia(1), "Otro", "", 99], [dia(1), "Orden de pago", "PAGO - 1", -1500]])
    salida = tmp_path / "salida" / "BBVA"
    assert cruzar_banco(entrada, salida, "2026-09")
    libro = openpyxl.load_workbook(salida / "cruce - BBVA 2026-09.xlsx")
    assert libro.sheetnames == ["Resumen", "Banco", "Sistema"]
    banco = list(libro["Banco"].iter_rows(values_only=True))
    sistema = list(libro["Sistema"].iter_rows(values_only=True))
    assert banco[0] == ("Fecha", "Código", "Descripción", "Monto", "Estado", "Fila en Sistema", "Criterio")
    assert sistema[0][5] == "Fila en Banco"
    # La transferencia (fila 3 del banco) cruzó con la orden de pago (fila 3 del sistema).
    assert [fila[4:] for fila in banco[1:]] == [("Sin cruzar", None, None), ("Cruzado", 3, "Exacto")]
    assert [fila[4:] for fila in sistema[1:]] == [("Sin cruzar", None, None), ("Cruzado", 3, "Exacto")]


def test_si_falla_no_escribe_y_borra_el_resultado_anterior(tmp_path, capsys):
    entrada = carpeta_banco(tmp_path / "entrada", [[dia(1), "A", "", 1]], [["mal", "A", "", 1]])
    salida = tmp_path / "salida" / "BBVA"
    salida.mkdir(parents=True)
    (salida / "cruce - BBVA 2026-09.xlsx").write_text("viejo")
    assert not cruzar_banco(entrada, salida, "2026-09")
    assert not (salida / "cruce - BBVA 2026-09.xlsx").exists()
    assert "la Fecha no es una fecha" in capsys.readouterr().out
